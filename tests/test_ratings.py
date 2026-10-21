import io
from datetime import date
from decimal import Decimal

from django.core.cache import cache
from django.test import TestCase
from django.urls import reverse
from openpyxl import load_workbook

from accounts.models import Role
from achievements.models import Achievement, Criterion, ManualScore
from core.models import SiteSettings
from ratings.models import StudentScore
from ratings.selection import Period, rating_rows
from ratings.services import recalculate, recalculate_semester
from reports.exports import build_excel

from .factories import make_staff, make_structure, make_student


class RatingBase(TestCase):
    def setUp(self):
        cache.clear()
        self.s = make_structure()
        self.a = make_student("309251100011", self.s["g1"], gpa="5.00")
        self.b = make_student("309251100012", self.s["g1"], gpa="5.00")
        self.c = make_student("309251100013", self.s["g2"], gpa="2.50")
        self.c.last_name = "=HYPERLINK(\"x\")"
        self.c.save()
        self.sport = Criterion.objects.get(code="sport")
        self.boss = make_staff("rahbar1", Role.LEADERSHIP)
        self.dean = make_staff("dekan1", Role.FACULTY_ADMIN, faculty=self.s["f1"])
        self.sem = self.s["sem1"]
        self.approve(self.a, "3")
        for st in (self.a, self.b, self.c):
            recalculate(st, self.sem)

    def approve(self, student, pts, sem=None):
        return Achievement.objects.create(
            student=student, semester=sem or self.sem, criterion=self.sport, title="Turnir", event_date=date(2026, 10, 1),
            status=Achievement.Status.APPROVED, points=Decimal(pts),
        )

    def period(self):
        return Period(f"s{self.sem.pk}", "1-semestr", semester=self.sem, year=self.s["year"])


class SelectionTests(RatingBase):
    def test_ranking_with_ties(self):
        rows = rating_rows(self.boss, self.period())
        self.assertEqual([r.student for r in rows], [self.a, self.b, self.c])
        self.assertEqual([r.rank for r in rows], [1, 2, 3])
        self.approve(self.b, "3")
        recalculate(self.b, self.sem)
        rows = rating_rows(self.boss, self.period())
        self.assertEqual([r.rank for r in rows], [1, 1, 3])

    def test_scope_limits_rows(self):
        rows = rating_rows(self.dean, self.period())
        self.assertEqual({r.student for r in rows}, {self.a, self.b})

    def test_year_is_semester_average(self):
        sem2 = self.s["sem2"]
        StudentScore.objects.create(student=self.a, semester=sem2, total=Decimal("0"), breakdown={"sport": "0"})
        year = Period(f"y{self.s['year'].pk}", "yil", year=self.s["year"])
        row = next(r for r in rating_rows(self.boss, year) if r.student == self.a)
        first = StudentScore.objects.get(student=self.a, semester=self.sem).total
        self.assertEqual(row.total, (first / 2).quantize(Decimal("0.01")))
        self.assertEqual(Decimal(row.breakdown["sport"]), Decimal("1.50"))

    def test_bulk_recalculation_matches_single(self):
        ManualScore.objects.create(student=self.b, semester=self.sem, criterion=Criterion.objects.get(code="odob"), points=Decimal("4"))
        self.approve(self.c, "5")
        expected = {}
        for st in (self.a, self.b, self.c):
            expected[st.pk] = recalculate(st, self.sem).breakdown
        StudentScore.objects.update(total=0, breakdown={})
        recalculate_semester(self.sem)
        for st in (self.a, self.b, self.c):
            self.assertEqual(StudentScore.objects.get(student=st, semester=self.sem).breakdown, expected[st.pk])


class RatingViewTests(RatingBase):
    def test_student_blocked_from_rating_but_sees_top10(self):
        self.client.force_login(self.a)
        self.assertEqual(self.client.get(reverse("ratings:rating")).status_code, 403)
        self.assertEqual(self.client.get(reverse("ratings:analytics")).status_code, 403)
        r = self.client.get(reverse("ratings:top10"))
        self.assertContains(r, "Sizning o‘rningiz: 1")

    def test_dean_rating_scoped(self):
        self.client.force_login(self.dean)
        r = self.client.get(reverse("ratings:rating"))
        self.assertContains(r, self.a.full_name)
        self.assertNotContains(r, "HYPERLINK")

    def test_dean_cannot_open_foreign_student_card(self):
        self.client.force_login(self.dean)
        self.assertEqual(self.client.get(reverse("ratings:student", args=[self.c.pk])).status_code, 404)
        self.assertEqual(self.client.get(reverse("ratings:student", args=[self.a.pk])).status_code, 200)

    def test_excel_export(self):
        self.client.force_login(self.boss)
        r = self.client.get(reverse("ratings:export", args=["excel"]))
        self.assertEqual(r.status_code, 200)
        ws = load_workbook(io.BytesIO(r.content)).active
        self.assertEqual(ws["B5"].value, self.a.full_name)
        names = [ws.cell(row=i, column=2).value for i in range(5, 8)]
        self.assertTrue(any(n.startswith("'=") for n in names), names)  # formula injection'dan himoya

    def test_pdf_export(self):
        self.client.force_login(self.boss)
        r = self.client.get(reverse("ratings:export", args=["pdf"]))
        self.assertEqual(r.status_code, 200)
        self.assertTrue(r.content.startswith(b"%PDF"))

    def test_unknown_export_format(self):
        self.client.force_login(self.boss)
        self.assertEqual(self.client.get(reverse("ratings:export", args=["csv"])).status_code, 404)

    def test_analytics_renders(self):
        for user in (self.boss, self.dean):
            self.client.force_login(user)
            r = self.client.get(reverse("ratings:analytics"))
            self.assertContains(r, "Mezonlar bo‘yicha bajarilish")

    def test_build_excel_empty(self):
        data = build_excel([], list(Criterion.objects.all()), {"title": "T", "period": "P", "scope": "S", "university": "U"})
        self.assertTrue(data.startswith(b"PK"))


class PublicApiTests(RatingBase):
    def test_public_stats(self):
        r = self.client.get(reverse("ratings:public_stats"))
        self.assertEqual(r.status_code, 200)
        data = r.json()
        self.assertEqual(data["students_rated"], 3)
        self.assertEqual(data["top10"][0]["name"], "Valiyev A.")
        self.assertNotIn("hemis", r.content.decode().lower())
        self.assertEqual(r["Access-Control-Allow-Origin"], "*")

    def test_public_stats_disabled(self):
        st = SiteSettings.load()
        st.public_stats_enabled = False
        st.save()
        self.assertEqual(self.client.get(reverse("ratings:public_stats")).status_code, 404)


class CriteriaEditorTests(RatingBase):
    def payload(self, **changes):
        qs = list(Criterion.objects.order_by("order", "id"))
        data = {"form-TOTAL_FORMS": len(qs), "form-INITIAL_FORMS": len(qs)}
        for i, c in enumerate(qs):
            data.update({f"form-{i}-id": c.pk, f"form-{i}-short_title": c.short_title, f"form-{i}-max_points": c.max_points,
                         f"form-{i}-order": c.order, f"form-{i}-is_active": "on"})
            for k, v in changes.get(c.code, {}).items():
                if v is None:
                    data.pop(f"form-{i}-{k}")
                else:
                    data[f"form-{i}-{k}"] = v
        return data

    def test_only_leadership(self):
        self.client.force_login(self.dean)
        self.assertEqual(self.client.get(reverse("achievements:criteria")).status_code, 403)

    def test_sum_must_be_100(self):
        self.client.force_login(self.boss)
        r = self.client.post(reverse("achievements:criteria"), self.payload(sport={"max_points": 6}))
        self.assertContains(r, "100 bo‘lishi kerak")
        self.assertEqual(Criterion.objects.get(code="sport").max_points, 5)

    def test_change_recalculates_current_semester(self):
        self.client.force_login(self.boss)
        # sport 5 -> 2, boshqa 5 -> 8: yig'indi 100 qoladi; talaba A ning 3 balli sporti 2 ga tushadi.
        r = self.client.post(reverse("achievements:criteria"), self.payload(sport={"max_points": 2}, boshqa={"max_points": 8}))
        self.assertRedirects(r, reverse("achievements:criteria"))
        score = StudentScore.objects.get(student=self.a, semester=self.sem)
        self.assertEqual(Decimal(score.breakdown["sport"]), Decimal("2"))

    def test_level_points_capped(self):
        self.client.force_login(self.boss)
        vol = Criterion.objects.get(code="volontyorlik")
        levels = list(vol.levels.all())
        data = {"title": vol.title, "short_title": vol.short_title, "description": "",
                "levels-TOTAL_FORMS": len(levels), "levels-INITIAL_FORMS": len(levels), "levels-MIN_NUM_FORMS": 0, "levels-MAX_NUM_FORMS": 1000}
        for i, lv in enumerate(levels):
            data.update({f"levels-{i}-id": lv.pk, f"levels-{i}-title": lv.title, f"levels-{i}-points": "9", f"levels-{i}-order": lv.order})
        r = self.client.post(reverse("achievements:criterion_edit", args=[vol.pk]), data)
        self.assertContains(r, "maksimumidan")
