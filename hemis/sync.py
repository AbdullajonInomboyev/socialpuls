"""HEMIS dan ommaviy sinxronizatsiya va Excel orqali talabalar importi.

Ikkalasi ham HemisSyncRun yozuvini qoldiradi; moderator sahifasi va `hemis_sync` buyrug'i shu funksiyalarni chaqiradi.
"""
import logging

from django.core.cache import cache
from django.db import transaction
from django.utils import timezone

from audit.services import log_action
from ratings.services import recalculate_semester
from structure.models import Faculty, Group, Semester, Specialty

from .client import HemisError, get_data_client
from .models import HemisSyncRun
from .services import _parse_gpa, upsert_faculty, upsert_group, upsert_specialty, upsert_student

logger = logging.getLogger("socialpuls.hemis")
LOCK_KEY = "hemis-sync-lock"
FACULTY_STRUCTURE_CODE = "11"  # HEMIS: structureType 11 = fakultet


class SyncBusy(Exception):
    pass


def _finish(run, status, error=""):
    run.status = status
    run.error = error[:2000]
    run.finished_at = timezone.now()
    run.save()


def sync_all(user=None, request=None, students=True):
    """Fakultet, yo'nalish, guruh va (ixtiyoriy) talabalarni sinxronlaydi. Bir vaqtda faqat bitta jarayon."""
    if not cache.add(LOCK_KEY, 1, timeout=60 * 60):
        raise SyncBusy("Sinxronizatsiya allaqachon bajarilmoqda.")
    run = HemisSyncRun.objects.create(kind=HemisSyncRun.Kind.STUDENTS if students else HemisSyncRun.Kind.GROUPS, triggered_by=user)
    try:
        client = get_data_client()
        created = updated = 0
        faculties = {}
        for item in client.iter_items("data/department-list"):
            code = str((item.get("structureType") or {}).get("code") or "")
            if code and code != FACULTY_STRUCTURE_CODE:
                continue
            existed = Faculty.objects.filter(hemis_id=item.get("id")).exists()
            fac = upsert_faculty(item)
            if fac:
                faculties[fac.hemis_id] = fac
                created, updated = created + (not existed), updated + existed
        specialties = {}
        for item in client.iter_items("data/specialty-list"):
            fac = faculties.get((item.get("department") or {}).get("id")) or Faculty.objects.filter(
                hemis_id=(item.get("department") or {}).get("id")).first()
            sp = upsert_specialty(item, fac)
            if sp:
                specialties[sp.hemis_id] = sp
        for item in client.iter_items("data/group-list"):
            fac = faculties.get((item.get("department") or {}).get("id")) or Faculty.objects.filter(
                hemis_id=(item.get("department") or {}).get("id")).first()
            sp = specialties.get((item.get("specialty") or {}).get("id"))
            upsert_group(item, fac, sp)
        if students:
            for item in client.iter_items("data/student-list"):
                try:
                    _, was_created = upsert_student(item, recalc=False)
                except (ValueError, PermissionError) as exc:
                    logger.warning("Talaba o'tkazib yuborildi (%s): %s", item.get("student_id_number"), exc)
                    continue
                created, updated = created + was_created, updated + (not was_created)
            sem = Semester.current()
            if sem and not sem.is_closed:
                recalculate_semester(sem)
        run.created_count, run.updated_count = created, updated
        _finish(run, HemisSyncRun.Status.SUCCESS)
        log_action(request, "hemis.sync", f"HEMIS sinxronizatsiya: {created} yangi, {updated} yangilandi", actor=user)
    except HemisError as exc:
        _finish(run, HemisSyncRun.Status.FAILED, str(exc))
        log_action(request, "hemis.sync_failed", f"HEMIS sinxronizatsiya xatosi: {exc}", actor=user)
    except Exception as exc:
        logger.exception("HEMIS sinxronizatsiyada kutilmagan xato")
        _finish(run, HemisSyncRun.Status.FAILED, f"Kutilmagan xato: {exc}")
        raise
    finally:
        cache.delete(LOCK_KEY)
    return run


# ---------- Excel import ----------

EXCEL_COLUMNS = ["HEMIS ID", "Familiya", "Ism", "Otasining ismi", "Fakultet", "Yo‘nalish", "Guruh", "Kurs", "GPA"]


def excel_template():
    import io

    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = "Talabalar"
    ws.append(EXCEL_COLUMNS)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="C25700")
        cell.fill = PatternFill("solid", fgColor="FFF2E5")
    ws.append(["309251100001", "Valiyev", "Ali", "Vali o‘g‘li", "Iqtisodiyot fakulteti", "Iqtisodiyot", "IQT-AU-25", 1, 4.25])
    for col, w in zip("ABCDEFGHI", (16, 18, 16, 18, 28, 28, 14, 7, 7)):
        ws.column_dimensions[col].width = w
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _cell(v):
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v).strip()


def parse_excel(file, max_rows=20000):
    """Qatorlarni tekshiradi: (to'g'ri qatorlar ro'yxati, xatolar ro'yxati)."""
    from openpyxl import load_workbook

    try:
        wb = load_workbook(file, read_only=True, data_only=True)
    except Exception as exc:  # noqa: BLE001 - buzilgan fayl turli xato beradi
        return [], [f"Faylni o‘qib bo‘lmadi: {exc.__class__.__name__}. .xlsx formatida saqlang."]
    ws = wb.worksheets[0]
    rows = ws.iter_rows(values_only=True)
    header = [_cell(h).lower() for h in next(rows, [])]
    expected = [c.lower() for c in EXCEL_COLUMNS]
    if header[: len(expected)] != expected:
        return [], ["Ustunlar shablonga mos emas. Shablonni yuklab olib, o‘sha tartibda to‘ldiring."]
    good, errors, seen = [], [], set()
    for i, raw in enumerate(rows, start=2):
        if i - 1 > max_rows:
            errors.append(f"{max_rows} qatordan ko‘pi qabul qilinmaydi.")
            break
        vals = [_cell(v) for v in (list(raw) + [None] * 9)[:9]]
        if not any(vals):
            continue
        sid, last, first, middle, fac, spec, grp, course, gpa = vals
        problems = []
        if not sid.isdigit() or not 6 <= len(sid) <= 20:
            problems.append("HEMIS ID raqamlardan iborat bo‘lishi kerak")
        elif sid in seen:
            problems.append("HEMIS ID takrorlangan")
        if not last or not first:
            problems.append("familiya va ism majburiy")
        if not fac or not grp:
            problems.append("fakultet va guruh majburiy")
        if course and (not course.isdigit() or not 1 <= int(course) <= 7):
            problems.append("kurs 1–7 oralig‘ida")
        gpa_val = _parse_gpa(gpa) if gpa else None
        if gpa and gpa_val is None:
            problems.append("GPA 0–5 oralig‘ida")
        if problems:
            errors.append(f"{i}-qator: " + ", ".join(problems))
            continue
        seen.add(sid)
        good.append({"sid": sid, "last": last, "first": first, "middle": middle, "faculty": fac, "specialty": spec,
                     "group": grp, "course": int(course) if course else None, "gpa": gpa_val})
    return good, errors


@transaction.atomic
def import_rows(rows, user=None, request=None):
    """Tekshirilgan qatorlarni yozadi. Fakultet va guruh nomi bo'yicha topiladi yoki yaratiladi."""
    from accounts.models import Role, StudentProfile, User

    from .services import _cap

    run = HemisSyncRun.objects.create(kind=HemisSyncRun.Kind.EXCEL, triggered_by=user)
    created = updated = skipped = 0
    fac_cache, sp_cache, grp_cache = {}, {}, {}
    now = timezone.now()
    for r in rows:
        fac = fac_cache.get(r["faculty"].lower())
        if fac is None:
            fac = Faculty.objects.filter(name__iexact=r["faculty"]).first() or Faculty.objects.create(name=r["faculty"])
            fac_cache[r["faculty"].lower()] = fac
        sp = None
        if r["specialty"]:
            key = (fac.pk, r["specialty"].lower())
            sp = sp_cache.get(key)
            if sp is None:
                sp = Specialty.objects.filter(faculty=fac, name__iexact=r["specialty"]).first() or Specialty.objects.create(faculty=fac, name=r["specialty"])
                sp_cache[key] = sp
        key = (fac.pk, r["group"].lower())
        grp = grp_cache.get(key)
        if grp is None:
            grp = Group.objects.filter(faculty=fac, name__iexact=r["group"]).first() or Group.objects.create(faculty=fac, specialty=sp, name=r["group"])
            grp_cache[key] = grp

        u = User.objects.filter(hemis_id=r["sid"]).first()
        if u is not None and u.role != Role.STUDENT:
            skipped += 1
            continue
        is_new = u is None
        if is_new:
            u = User(username=f"s{r['sid']}", hemis_id=r["sid"], role=Role.STUDENT)
            u.set_unusable_password()
        u.last_name, u.first_name, u.middle_name = _cap(r["last"]), _cap(r["first"]), _cap(r["middle"])
        u.save()
        defaults = {"student_id_number": r["sid"], "faculty": fac, "specialty": sp, "group": grp, "course": r["course"], "hemis_synced_at": now}
        if r["gpa"] is not None:
            defaults.update(gpa=r["gpa"], gpa_synced_at=now)
        StudentProfile.objects.update_or_create(user=u, defaults=defaults)
        created, updated = created + is_new, updated + (not is_new)
    sem = Semester.current()
    if sem and not sem.is_closed:
        recalculate_semester(sem)
    run.created_count, run.updated_count = created, updated
    _finish(run, HemisSyncRun.Status.SUCCESS, f"{skipped} ta qator xodim hisobiga to‘g‘ri kelgani uchun o‘tkazib yuborildi." if skipped else "")
    log_action(request, "students.import", f"Excel import: {created} yangi, {updated} yangilandi, {skipped} o‘tkazildi", actor=user)
    return run


