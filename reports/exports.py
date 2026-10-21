"""Reytingni Excel (openpyxl) va PDF (reportlab) ko'rinishida tayyorlash."""
import io
from pathlib import Path

from django.utils import timezone

FONT_DIR = Path(__file__).resolve().parent / "fonts"
PDF_ROW_LIMIT = 3000


def _safe(value):
    """Excel formula injection'dan himoya: =, +, -, @ bilan boshlangan matn formula bo'lib qolmasin."""
    if isinstance(value, str) and value[:1] in ("=", "+", "-", "@", "\t", "\r"):
        return "'" + value
    return value


def build_excel(rows, criteria, meta):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "Reyting"
    headers = ["O‘rin", "F.I.Sh.", "HEMIS ID", "Fakultet", "Guruh", "Kurs"] + [c.short_title for c in criteria] + ["Jami", "Tasdiqlangan yutuqlar"]
    last_col = get_column_letter(len(headers))

    ws["A1"] = meta["title"]
    ws["A1"].font = Font(size=14, bold=True)
    ws.merge_cells(f"A1:{last_col}1")
    ws["A2"] = f"{meta['period']} · {meta['scope']} · Tayyorlangan: {timezone.localtime():%d.%m.%Y %H:%M}"
    ws["A2"].font = Font(size=10, color="64748B")
    ws.merge_cells(f"A2:{last_col}2")

    head_fill = PatternFill("solid", fgColor="FFF2E5")
    thin = Side(style="thin", color="E8EDF4")
    for col, text in enumerate(headers, start=1):
        cell = ws.cell(row=4, column=col, value=text)
        cell.font = Font(bold=True, color="C25700")
        cell.fill = head_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(bottom=thin)
    ws.row_dimensions[4].height = 34

    for i, r in enumerate(rows, start=5):
        p = r.profile
        values = [
            r.rank, r.student.full_name, r.student.hemis_id or "",
            p.faculty.name if p and p.faculty else "", p.group.name if p and p.group else "", p.course if p else None,
        ] + [float(v) for v in r.cells] + [float(r.total), r.approved]
        for col, v in enumerate(values, start=1):
            cell = ws.cell(row=i, column=col, value=_safe(v))
            if 7 <= col <= 6 + len(criteria) + 1:
                cell.number_format = "0.00"
        ws.cell(row=i, column=7 + len(criteria)).font = Font(bold=True)

    widths = [7, 34, 15, 28, 14, 6] + [11] * len(criteria) + [9, 12]
    for col, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(col)].width = w
    ws.freeze_panes = "C5"
    if rows:
        ws.auto_filter.ref = f"A4:{last_col}{4 + len(rows)}"

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


_fonts_ready = False


def _register_fonts():
    global _fonts_ready
    if _fonts_ready:
        return
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont

    pdfmetrics.registerFont(TTFont("Outfit", str(FONT_DIR / "Outfit-Regular.ttf")))
    pdfmetrics.registerFont(TTFont("Outfit-Bold", str(FONT_DIR / "Outfit-Bold.ttf")))
    _fonts_ready = True


def _pdf_text(value):
    # Outfit shriftida U+02BB/02BC yo'q; HEMIS ismlarida uchraydi, ko'rinishi bir xil belgiga almashtiriladi.
    return str(value).replace("ʻ", "‘").replace("ʼ", "’").replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


def build_pdf(rows, criteria, meta):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    _register_fonts()
    truncated = len(rows) > PDF_ROW_LIMIT
    rows = rows[:PDF_ROW_LIMIT]
    orange, ink, muted, line, soft = (colors.HexColor(h) for h in ("#C25700", "#1E293B", "#64748B", "#E8EDF4", "#FFF2E5"))
    st_title = ParagraphStyle("t", fontName="Outfit-Bold", fontSize=15, leading=19, textColor=ink)
    st_meta = ParagraphStyle("m", fontName="Outfit", fontSize=8.5, leading=11, textColor=muted)
    st_head = ParagraphStyle("h", fontName="Outfit-Bold", fontSize=6, leading=7, textColor=orange, alignment=1, splitLongWords=0)
    st_cell = ParagraphStyle("c", fontName="Outfit", fontSize=7.5, leading=9, textColor=ink)

    buf = io.BytesIO()
    generated = timezone.localtime()

    def footer(canvas, doc):
        canvas.saveState()
        canvas.setFont("Outfit", 7)
        canvas.setFillColor(muted)
        canvas.drawString(12 * mm, 8 * mm, f"{meta['university']} · Ijtimoiy Faollik · {generated:%d.%m.%Y %H:%M}")
        canvas.drawRightString(landscape(A4)[0] - 12 * mm, 8 * mm, f"{doc.page}-sahifa")
        canvas.restoreState()

    doc = SimpleDocTemplate(
        buf, pagesize=landscape(A4), leftMargin=12 * mm, rightMargin=12 * mm, topMargin=12 * mm, bottomMargin=14 * mm,
        title=meta["title"], author=meta["university"],
    )
    story = [
        Paragraph(_pdf_text(meta["title"]), st_title),
        Spacer(1, 2 * mm),
        Paragraph(_pdf_text(f"{meta['period']} · {meta['scope']} · {len(rows)} talaba"), st_meta),
    ]
    if truncated:
        story.append(Paragraph(f"PDF faqat birinchi {PDF_ROW_LIMIT} qatorni o‘z ichiga oladi. To‘liq ro‘yxat uchun Excel’dan foydalaning.", st_meta))
    story.append(Spacer(1, 4 * mm))

    head = [Paragraph(t, st_head) for t in ["O‘rin", "F.I.Sh.", "Guruh", "Kurs"]]
    head += [Paragraph(_pdf_text(f"{c.short_title} ({c.max_points})"), st_head) for c in criteria]
    head.append(Paragraph("Jami", st_head))
    data = [head]
    for r in rows:
        p = r.profile
        data.append(
            [str(r.rank), Paragraph(_pdf_text(r.student.full_name), st_cell), _pdf_text(p.group.name if p and p.group else ""),
             str(p.course or "") if p else ""]
            + [_num(v) for v in r.cells]
            + [_num(r.total)]
        )

    crit_w = 41
    widths = [32, 150, 62, 28] + [crit_w] * len(criteria) + [40]
    available = landscape(A4)[0] - 24 * mm
    widths[1] = max(110, available - (sum(widths) - widths[1]))
    table = Table(data, colWidths=widths, repeatRows=1)
    style = [
        ("FONT", (0, 1), (-1, -1), "Outfit", 7.5),
        ("FONT", (-1, 1), (-1, -1), "Outfit-Bold", 8),
        ("BACKGROUND", (0, 0), (-1, 0), soft),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("ALIGN", (0, 1), (0, -1), "CENTER"),
        ("ALIGN", (3, 1), (-1, -1), "CENTER"),
        ("LINEBELOW", (0, 0), (-1, -1), 0.4, line),
        ("TOPPADDING", (0, 0), (-1, -1), 3),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
    ]
    for i in range(2, len(data), 2):
        style.append(("BACKGROUND", (0, i), (-1, i), colors.HexColor("#F8FAFC")))
    table.setStyle(TableStyle(style))
    story.append(table)
    doc.build(story, onFirstPage=footer, onLaterPages=footer)
    return buf.getvalue()


def _num(value):
    s = f"{value:.2f}".rstrip("0").rstrip(".")
    return s.replace(".", ",")
